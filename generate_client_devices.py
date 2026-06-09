import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import random
from datetime import datetime

# Create a new workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Client Devices"

# Define styles
header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=12)
client_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
client_font = Font(bold=True, size=11)
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Set column widths
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 25
ws.column_dimensions['D'].width = 15

# Add title
ws['A1'] = "MEGHA SMART - CLIENT DEVICE ALLOCATION"
ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
ws['A1'].fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
ws.merge_cells('A1:D1')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 25

# Add headers
headers = ['S.No', 'Client Name', 'Device ID', 'Status']
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col)
    cell.value = header
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border

ws.row_dimensions[3].height = 20

# Client data
clients = [
    {'name': 'Rajesh Kumar', 'email': 'rajesh.kumar@meghasmart.com', 'flatNo': 'A1-101'},
    {'name': 'Priya Singh', 'email': 'priya.singh@meghasmart.com', 'flatNo': 'A2-202'},
    {'name': 'Arjun Patel', 'email': 'arjun.patel@meghasmart.com', 'flatNo': 'A3-303'}
]

# Generate device IDs
year = datetime.now().year
device_ids = [f"GVM{year}{str(i).zfill(5)}" for i in range(1, 91)]  # 90 devices total (30 per client)
random.shuffle(device_ids)  # Shuffle to randomly assign

# Fill in data
row = 4
serial_no = 1
device_index = 0

for client_idx, client in enumerate(clients):
    # Add client header row
    ws.merge_cells(f'A{row}:D{row}')
    client_cell = ws[f'A{row}']
    client_cell.value = f"Client {client_idx + 1}: {client['name']} ({client['email']}) - Flat: {client['flatNo']}"
    client_cell.fill = client_fill
    client_cell.font = client_font
    client_cell.alignment = Alignment(horizontal='left', vertical='center')
    client_cell.border = border
    ws.row_dimensions[row].height = 18
    row += 1
    
    # Add 30 devices for this client
    for device_count in range(30):
        ws.cell(row=row, column=1).value = serial_no
        ws.cell(row=row, column=2).value = client['name']
        ws.cell(row=row, column=3).value = device_ids[device_index]
        ws.cell(row=row, column=4).value = "Active"
        
        # Apply borders and alignment
        for col in range(1, 5):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            if col in [1, 4]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        serial_no += 1
        device_index += 1
        row += 1
    
    # Add blank row between clients
    row += 1

# Add summary section
summary_row = row + 1
ws[f'A{summary_row}'] = "SUMMARY"
ws[f'A{summary_row}'].font = Font(bold=True, size=12, color="FFFFFF")
ws[f'A{summary_row}'].fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
ws.merge_cells(f'A{summary_row}:D{summary_row}')

summary_row += 1
ws[f'A{summary_row}'] = "Total Clients"
ws[f'B{summary_row}'] = len(clients)
ws[f'A{summary_row}'].font = Font(bold=True)

summary_row += 1
ws[f'A{summary_row}'] = "Devices per Client"
ws[f'B{summary_row}'] = 30
ws[f'A{summary_row}'].font = Font(bold=True)

summary_row += 1
ws[f'A{summary_row}'] = "Total Devices"
ws[f'B{summary_row}'] = len(clients) * 30
ws[f'A{summary_row}'].font = Font(bold=True)

# Save the workbook
output_file = "Megha_Smart_Client_Devices.xlsx"
wb.save(output_file)
print(f"✅ Excel file created successfully: {output_file}")
print(f"\n📊 Summary:")
print(f"   - Total Clients: {len(clients)}")
print(f"   - Devices per Client: 30")
print(f"   - Total Devices: {len(clients) * 30}")
print(f"\n📁 File saved at: {output_file}")
